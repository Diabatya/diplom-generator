"""Generate test assets into testdata/: template.pptx, template.pdf,
list.xlsx (10 rows, duplicates), list_bad.xlsx (empty cell),
list_long.xlsx (very long name)."""

import io
import os

try:
    import pymupdf as fitz
except ImportError:
    import fitz
import openpyxl
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Cm, Pt

from app.core.fonts import find_font_file

OUT = os.path.join(os.path.dirname(__file__), "testdata")
FIO_PH = "@@FIO@@"
REG_PH = "@@REGION@@"


def _logo_png() -> bytes:
    doc = fitz.open()
    page = doc.new_page(width=200, height=200)
    page.draw_circle(fitz.Point(100, 100), 80, color=(0.8, 0.6, 0.1),
                     fill=(0.95, 0.8, 0.3), width=6)
    page.draw_circle(fitz.Point(100, 100), 40, color=(0.8, 0.6, 0.1),
                     fill=(0.9, 0.7, 0.2), width=3)
    pix = page.get_pixmap(dpi=96)
    data = pix.tobytes("png")
    doc.close()
    return data


def make_pptx(path):
    prs = Presentation()
    prs.slide_width = Cm(29.7)
    prs.slide_height = Cm(21.0)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # logo picture — needed to verify combined mode keeps images
    slide.shapes.add_picture(io.BytesIO(_logo_png()), Cm(25.5), Cm(15), Cm(3), Cm(3))

    bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)  # rect
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0xF5, 0xF0, 0xE0)
    bg.line.fill.background()
    bg.shadow.inherit = False

    band = slide.shapes.add_shape(1, 0, Cm(2), prs.slide_width, Cm(3))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(0x1F, 0x3F, 0x7F)
    band.line.fill.background()
    band.shadow.inherit = False

    tb = slide.shapes.add_textbox(Cm(3), Cm(2.4), Cm(23.7), Cm(2.2))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = "ДИПЛОМ"
    r.font.size = Pt(44)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    fio = slide.shapes.add_textbox(Cm(3), Cm(9), Cm(23.7), Cm(2.5))
    fr = fio.text_frame.paragraphs[0].add_run()
    fr.text = FIO_PH
    fr.font.size = Pt(40)
    fr.font.name = "Arial"
    fr.font.bold = True
    fr.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    fio.text_frame.paragraphs[0].alignment = 2  # center

    reg = slide.shapes.add_textbox(Cm(3), Cm(12.5), Cm(23.7), Cm(1.5))
    rr = reg.text_frame.paragraphs[0].add_run()
    rr.text = REG_PH
    rr.font.size = Pt(24)
    rr.font.name = "Arial"
    rr.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    # second FIO occurrence — small line at the bottom
    fio2 = slide.shapes.add_textbox(Cm(3), Cm(18.5), Cm(23.7), Cm(1.2))
    f2 = fio2.text_frame.paragraphs[0].add_run()
    f2.text = "Выдан: " + FIO_PH
    f2.font.size = Pt(14)
    f2.font.name = "Arial"
    f2.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    prs.save(path)


def make_pdf(path, font_path):
    doc = fitz.open()
    page = doc.new_page(width=842, height=595)  # A4 landscape
    page.draw_rect(fitz.Rect(20, 20, 822, 575), color=(0.12, 0.25, 0.5), width=3)
    page.draw_rect(fitz.Rect(28, 28, 814, 567), color=(0.12, 0.25, 0.5), width=1)
    kw = dict(fontname="UArial", fontfile=font_path)
    page.insert_textbox(
        fitz.Rect(60, 50, 782, 140), "ДИПЛОМ", fontsize=48,
        color=(0.12, 0.25, 0.5), align=fitz.TEXT_ALIGN_CENTER, **kw,
    )
    page.insert_textbox(
        fitz.Rect(60, 240, 782, 320), FIO_PH, fontsize=36,
        color=(0, 0, 0), align=fitz.TEXT_ALIGN_CENTER, **kw,
    )
    page.insert_textbox(
        fitz.Rect(60, 330, 782, 380), REG_PH, fontsize=22,
        color=(0.3, 0.3, 0.3), align=fitz.TEXT_ALIGN_CENTER, **kw,
    )
    doc.save(path)
    doc.close()


NAMES = [
    "Иванов Иван Иванович", "Петрова Анна Сергеевна", "Сидоров Алексей Петрович",
    "Кузнецова Мария Игоревна", "Иванов Иван Иванович", "Смирнов Олег Викторович",
    "Васильева Елена Дмитриевна", "Николаев Павел Андреевич", "Фёдорова Ольга Павловна",
    "Морозов Дмитрий Сергеевич",
]
REGIONS = ["Москва", "Санкт-Петербург", "Новосибирск", "Казань", "Екатеринбург",
           "Самара", "Ростов-на-Дону", "Уфа", "Красноярск", "Воронеж"]


def make_xlsx(path, names, regions, blank_cell=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ФИО", "Регион"])
    for i, (n, r) in enumerate(zip(names, regions)):
        ws.append(["" if (blank_cell and i == 3) else n, r])
    wb.save(path)


def main():
    os.makedirs(OUT, exist_ok=True)
    font_path = find_font_file("Arial") or find_font_file("DejaVu Sans")
    assert font_path, "no cyrillic TTF found on this machine"
    make_pptx(os.path.join(OUT, "template.pptx"))
    make_pdf(os.path.join(OUT, "template.pdf"), font_path)
    make_xlsx(os.path.join(OUT, "list.xlsx"), NAMES, REGIONS)
    make_xlsx(os.path.join(OUT, "list_bad.xlsx"), NAMES, REGIONS, blank_cell=True)
    make_xlsx(
        os.path.join(OUT, "list_long.xlsx"),
        ["Абдурахмангаджиев Константин Константинопольский-Оглы"] + NAMES[:3],
        REGIONS[:4],
    )
    print("testdata written to", OUT)
    print("pdf font:", font_path)


if __name__ == "__main__":
    main()
