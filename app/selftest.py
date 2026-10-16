"""--selftest: run core checks without opening the GUI.

Generates its own minimal fixtures in a temp dir (testdata/ is not shipped
in the exe), exercises excel/pptx/pdf/font paths, and exits non-zero on
any failure. Used by CI to prove the packaged onefile binary actually
starts and works on Windows.
"""

import io
import os
import platform
import sys
import tempfile
import traceback

_results = []


def _check(name, cond, extra=""):
    ok = bool(cond)
    _results.append((name, ok))
    print(f"[{'OK' if ok else 'FAIL'}] {name} {extra}")
    return ok


def _make_fixtures(tmp):
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    import openpyxl
    from pptx import Presentation
    from pptx.util import Cm, Pt

    xlsx = os.path.join(tmp, "list.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ФИО", "Регион"])
    for n, r in [("Иванов Иван Иванович", "Москва"),
                 ("Петрова Анна Сергеевна", "Казань")]:
        ws.append([n, r])
    wb.save(xlsx)

    pptx_path = os.path.join(tmp, "template.pptx")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(29.7), Cm(21)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Cm(3), Cm(9), Cm(23.7), Cm(2.5))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "@@FIO@@"
    run.font.size = Pt(40)
    reg = slide.shapes.add_textbox(Cm(3), Cm(12.5), Cm(23.7), Cm(1.5))
    rr = reg.text_frame.paragraphs[0].add_run()
    rr.text = "@@REGION@@"
    rr.font.size = Pt(24)
    prs.save(pptx_path)

    pdf_path = os.path.join(tmp, "template.pdf")
    doc = fitz.open()
    page = doc.new_page(width=842, height=595)
    page.draw_rect(fitz.Rect(20, 20, 822, 575), color=(0.12, 0.25, 0.5), width=2)
    page.insert_text(fitz.Point(60, 80), "DIPLOM @@FIO@@ @@REGION@@",
                     fontsize=20, color=(0, 0, 0))
    doc.save(pdf_path)
    doc.close()
    return xlsx, pptx_path, pdf_path


def main():
    print("=== DiplomGenerator --selftest ===")
    print("Python:", sys.version.split()[0])
    print("OS:", platform.platform())
    print("Frozen:", getattr(sys, "frozen", False))

    from app.core import excel_reader, pdf_engine, pptx_engine
    from app.core.fonts import find_font_file, list_system_fonts, system_font_dirs
    from app.core.generator import MODE_SEPARATE, generate
    from app.core.models import Field, PdfPlacement, PptxPlacement, TEMPLATE_PDF, TEMPLATE_PPTX

    # ---- system fonts --------------------------------------------------
    dirs = system_font_dirs()
    print("Font dirs:", dirs)
    fonts = list_system_fonts(cyrillic_only=True)
    _check("system fonts found", len(fonts) > 0, f"{len(fonts)} cyrillic-capable")
    names_l = [n.lower() for n, _ in fonts]
    for probe in ("arial", "calibri", "times"):
        hit = [p for n, p in fonts if probe in n.lower()]
        print(f"  font '{probe}':", hit[0] if hit else "not found")
    font_path = find_font_file("Arial") or (fonts[0][1] if fonts else None)
    _check("cyrillic TTF resolved", bool(font_path), str(font_path))

    # ---- font measurement / autofit ------------------------------------
    if font_path:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
        f = fitz.Font(fontfile=font_path)
        w = f.text_length("Иванов", fontsize=20)
        _check("font measurement (autofit)", w > 0, f"{w:.1f} pt wide")

    # ---- fixtures + excel ----------------------------------------------
    tmp = tempfile.mkdtemp(prefix="diplom_selftest_")
    xlsx, pptx_path, pdf_path = _make_fixtures(tmp)
    headers, rows = excel_reader.read_xlsx(xlsx)
    _check("excel read", headers == ["ФИО", "Регион"] and len(rows) == 2)
    cm = {"ФИО": 0, "Регион": 1}

    # ---- pptx generation -----------------------------------------------
    from pptx import Presentation
    shapes = {s["shape_id"]: s for s in pptx_engine.list_text_shapes(
        Presentation(pptx_path))}
    fields = [
        Field("ФИО", [PptxPlacement(0, sid, pattern=s["text"],
                                  token="@@FIO@@", mode="token")
                      for sid, s in shapes.items() if "@@FIO@@" in s["text"]]),
        Field("Регион", [PptxPlacement(0, sid, pattern=s["text"],
                                     token="@@REGION@@", mode="token")
                         for sid, s in shapes.items() if "@@REGION@@" in s["text"]]),
    ]
    rep = generate(pptx_path, TEMPLATE_PPTX, fields, cm, rows,
                   os.path.join(tmp, "pptx_out"), MODE_SEPARATE, 0, set())
    _check("pptx generation", len(rep["created"]) == 2,
           str([c[1] for c in rep["created"]]))

    # ---- facsimile signature -------------------------------------------
    from app.core.models import SignatureSpec
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore
    sig_png = os.path.join(tmp, "sig.png")
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 120, 50))
    pix.clear_with(30)
    pix.save(sig_png)
    spec = SignatureSpec(True, sig_png, (400, 480, 700, 545))
    rep_sig = generate(pptx_path, TEMPLATE_PPTX, fields, cm, rows,
                       os.path.join(tmp, "pptx_sig"), MODE_SEPARATE, 0, set(),
                       signature=spec)
    sig_pics = [s for s in Presentation(
        os.path.join(tmp, "pptx_sig", rep_sig["created"][0][1])
    ).slides[0].shapes
        if s.name == pptx_engine.SIGNATURE_PICTURE_NAME] \
        if rep_sig["created"] else []
    _check("pptx signature stamped", len(rep_sig["created"]) == 2
           and len(sig_pics) == 1, str(sig_pics))
    if rep["created"]:
        texts = [sh.text_frame.text
                 for sh in Presentation(
                     os.path.join(tmp, "pptx_out", rep["created"][0][1])
                 ).slides[0].shapes if sh.has_text_frame]
        _check("pptx text substituted",
               any("Иванов Иван Иванович" in t for t in texts)
               and not any("@@" in t for t in texts))

    # ---- pdf generation + preview render --------------------------------
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    doc = fitz.open(pdf_path)
    fio_rect = doc[0].search_for("@@FIO@@")[0]
    reg_rect = doc[0].search_for("@@REGION@@")[0]
    doc.close()
    grow = lambda r: (r.x0 - 30, r.y0 - 14, r.x1 + 30, r.y1 + 14)
    pfields = [
        Field("ФИО", [PdfPlacement(0, grow(fio_rect), font_path, 20,
                                   "#000000", "left")]),
        Field("Регион", [PdfPlacement(0, grow(reg_rect), font_path, 20,
                                      "#000000", "left")]),
    ]
    rep2 = generate(pdf_path, TEMPLATE_PDF, pfields, cm, rows,
                    os.path.join(tmp, "pdf_out"), MODE_SEPARATE, 0, set())
    _check("pdf generation", len(rep2["created"]) == 2,
           str([c[1] for c in rep2["created"]]))
    if rep2["created"]:
        out_pdf = os.path.join(tmp, "pdf_out", rep2["created"][0][1])
        d = fitz.open(out_pdf)
        txt = d[0].get_text()
        pix = d[0].get_pixmap(dpi=72)
        preview = os.path.join(tmp, "preview.png")
        pix.save(preview)
        d.close()
        _check("pdf text substituted",
               "Иванов" in txt.replace("\u00a0", " "))
        _check("pdf preview rendered", os.path.getsize(preview) > 1000)

    # ---- Qt plugins packaged (offscreen, no window shown) ---------------
    try:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication
        app = QApplication.instance() or QApplication(["selftest"])
        _check("Qt platform plugin loads (offscreen)", app is not None)
        app.quit()
    except Exception as e:
        _check("Qt platform plugin loads (offscreen)", False, repr(e))

    failed = [n for n, ok in _results if not ok]
    print(f"=== selftest: {len(_results) - len(failed)}/{len(_results)} checks OK ===")
    if failed:
        print("FAILED:", ", ".join(failed))
        return 1
    print("SELFTEST OK")
    return 0


class _Tee:
    """Mirror stdout/stderr into a log file (the packaged exe has no console)."""

    def __init__(self, fh):
        self.fh = fh
        self.out = sys.stdout
        self.err = sys.stderr

    def write(self, s):
        try:
            if self.out:
                self.out.write(s)
        except Exception:
            pass
        self.fh.write(s)

    def flush(self):
        try:
            if self.out:
                self.out.flush()
        except Exception:
            pass
        self.fh.flush()


def run():
    log_path = os.path.abspath("selftest_result.log")
    if "--selftest-log" in sys.argv:
        i = sys.argv.index("--selftest-log")
        if i + 1 < len(sys.argv):
            log_path = os.path.abspath(sys.argv[i + 1])
    # Best effort: attach to the parent console when launched from cmd/pwsh
    # (GUI-subsystem exe has none of its own).
    if sys.platform == "win32":
        try:
            import ctypes
            ctypes.windll.kernel32.AttachConsole(-1)  # ATTACH_PARENT_PROCESS
        except Exception:
            pass
    try:
        fh = open(log_path, "w", encoding="utf-8")
    except OSError:
        fh = open(os.path.join(tempfile.gettempdir(), "selftest_result.log"),
                  "w", encoding="utf-8")
        log_path = fh.name
    tee = _Tee(fh)
    sys.stdout = tee
    sys.stderr = tee
    try:
        try:
            code = main()
        except Exception:
            traceback.print_exc()
            code = 1
    finally:
        print("log:", log_path)
        fh.flush()
        fh.close()
    # os._exit: interpreter teardown after a Qt session can abort and mask
    # the real result code — skip it, the report is already written.
    os._exit(code)
