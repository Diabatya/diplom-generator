"""Read .xlsx lists. All values are returned as text.

Numbers stored as numbers are converted without a trailing ".0"; numeric
cells with a zero-padding number format like ``00000`` are zero-padded back
so leading zeros are preserved.
"""

from __future__ import annotations

import datetime
import re

import openpyxl

_ZERO_PAD_RE = re.compile(r"^0+$")


def _cell_to_text(cell) -> str:
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if isinstance(v, bool):
        return "Да" if v else "Нет"
    if isinstance(v, (int, float)):
        fmt = (cell.number_format or "").strip()
        if isinstance(v, float):
            if v.is_integer():
                v = int(v)
            else:
                return repr(v)
        if _ZERO_PAD_RE.match(fmt):
            return str(v).zfill(len(fmt))
        return str(v)
    if isinstance(v, datetime.datetime):
        return v.strftime("%d.%m.%Y") if (v.hour, v.minute, v.second) == (0, 0, 0) else str(v)
    if isinstance(v, datetime.date):
        return v.strftime("%d.%m.%Y")
    return str(v)


def sheet_names(path: str) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_xlsx(path: str, sheet_name: str | None = None):
    """Return (headers, rows) — lists of strings. Row 1 is the header row."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        grid: list[list[str]] = []
        for row in ws.iter_rows():
            grid.append([_cell_to_text(c) for c in row])
        # drop fully empty trailing rows/cols
        while grid and all(v == "" for v in grid[-1]):
            grid.pop()
        if not grid:
            return [], []
        width = max(len(r) for r in grid)
        grid = [r + [""] * (width - len(r)) for r in grid]
        headers = grid[0]
        rows = grid[1:]
        # drop trailing empty columns
        while headers and headers[-1] == "" and all(r[len(headers) - 1] == "" for r in rows):
            headers.pop()
            for r in rows:
                r.pop()
        return headers, rows
    finally:
        wb.close()


INVALID_FILENAME_CHARS = '\\/:*?"<>|'
_INVALID_RE = re.compile("[" + re.escape(INVALID_FILENAME_CHARS) + "]")


def validate_rows(
    rows: list[list[str]],
    col_indexes: list[int],
    filename_col: int | None = None,
    headers: list[str] | None = None,
):
    """Return list of (excel_row_number, problem) for mapped columns.

    Empty cells are checked in every mapped column; characters that are
    invalid in Windows file names are checked in the column used for
    naming output files. excel_row_number is 1-based including the header
    row (data starts at 2); row number 0 marks a column-level problem.
    """
    problems: list[tuple[int, str]] = []
    empty_cols = columns_all_empty(rows, col_indexes)
    for ci in sorted(empty_cols):
        title = headers[ci] if headers and ci < len(headers) else f"#{ci + 1}"
        problems.append(
            (0, f'столбец «{title}» пуст во всех строках — возможно, файл '
                'содержит формулы без сохранённых значений; откройте и '
                'пересохраните его в Excel')
        )
    for i, row in enumerate(rows):
        excel_row = i + 2
        for ci in col_indexes:
            if ci is None or ci < 0 or ci in empty_cols:
                continue
            value = row[ci] if ci < len(row) else ""
            if value.strip() == "":
                problems.append((excel_row, "пустая ячейка"))
        if filename_col is not None and 0 <= filename_col < len(row):
            value = row[filename_col]
            if _INVALID_RE.search(value):
                problems.append(
                    (excel_row, f'недопустимые символы для имени файла: "{value}"')
                )
    return problems


def columns_all_empty(rows: list[list[str]], col_indexes: list[int]) -> set[int]:
    """Columns where every single data row is empty — a sign of formulas
    without cached values (openpyxl data_only returns None for them)."""
    out = set()
    for ci in set(col_indexes):
        if ci is None or ci < 0:
            continue
        if rows and all((ci >= len(r) or r[ci].strip() == "") for r in rows):
            out.add(ci)
    return out
